from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from functools import partial
import gc
import hashlib
import io
import re
from typing import Any

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
from requests.adapters import HTTPAdapter
import streamlit as st
import urllib3
from urllib3.util import Retry

# Suppress insecure request warnings when falling back to verify=False on MOH servers
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# --- Streamlit Page Configuration ---
st.set_page_config(
    page_title="DHIS2 TB Tracker Data Exporter",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --- Constants & Program / OrgUnit Mappings ---
REGISTRATION_PROGRAM_ID = "UZF0HrTlps0"
DIAGNOSIS_PROGRAM_ID = "GvywHD6crky"
SURVEILLANCE_PROGRAM_ID = "Lt6P15ps7f6"
CONTACT_PROGRAM_ID = "cQsXTtAJ3HW"

PROGRAM_MAP: dict[str, str] = {
    "Registration & Screening": REGISTRATION_PROGRAM_ID,
    "Diagnostic Evaluation": DIAGNOSIS_PROGRAM_ID,
    "TB Surveillance": SURVEILLANCE_PROGRAM_ID,
    "TB Contact Tracing and TPT": CONTACT_PROGRAM_ID,
}

PROGRAM_ID_TO_LABEL: dict[str, str] = {v: k for k, v in PROGRAM_MAP.items()}

PROGRAM_CASCADE_ORDER: dict[str, int] = {
    REGISTRATION_PROGRAM_ID: 1,
    DIAGNOSIS_PROGRAM_ID: 2,
    SURVEILLANCE_PROGRAM_ID: 3,
    CONTACT_PROGRAM_ID: 4,
    "Registration & Screening": 1,
    "Diagnostic Evaluation": 2,
    "TB Surveillance": 3,
    "TB Contact Tracing and TPT": 4,
}

# Includes both Partner/YTP OUs and Township NTP OUs for different user accounts and roles
TOWNSHIP_OU_MAP: dict[str, list[str]] = {
    "12. YTP/Pyi Gyi Khin (Hlaing)": ["YGCPqlBamwD", "KqjORlUe8Yc"],
    "12. YTP/Pyi Gyi Khin (Kyeemyindaing)": ["dE0upZMc4xG", "rTTCKrLpxTB"],
    "12. YTP/Pyi Gyi Khin (Dagon Myothit (South))": ["mCgl8tnCaee", "XHz6CPxTAbR"],
    "12. YTP/Pyi Gyi Khin (Thingangyun)": ["lJ9AHubJTjg", "aBfPB9AwbF5"],
    "12. YTP/Pyi Gyi Khin (South Okkalapa)": ["KSRVZsryAKz", "OeZsFpNKLP5"],
    "12. YTP/Pyi Gyi Khin (Mayangone)": ["bctwifUcioE", "mPwLv1cjror"],
    "12. YTP/Pyi Gyi Khin (Shwepyithar)": ["Z7Qf4wgtHz3", "aMAEOgli6W8"],
}

# Required leading column order: Date, OrgUnit, UPI, Name, Age, Sex
PRIORITY_LEADING_COLUMNS: list[str] = [
    "enrollment_date",
    "Org Unit Name",
    "orgUnit_id",
    "Unique ID (UPI)",
    "GEN - Name",
    "Age",
    "GEN - Sex",
]

# YgnTBPro target columns ordered with Date, OrgUnit, UPI, Name, Age, Sex first
SELECTED_COLUMN_LIST: list[str] = [
    # 1. Leading Priority Columns (Date, OrgUnit, UPI, Name, Age, Sex)
    "enrollment_date",
    "Org Unit Name",
    "orgUnit_id",
    "Unique ID (UPI)",
    "GEN - Name",
    "Age",
    "GEN - Sex",
    # 2. Remaining Demographic & Administrative Columns
    "Nationality",
    "Home Address",
    "GEN - Contact phone number (local)",
    "NRC No.",
    "GEN - Date of birth",
    "GEN - Date of birth is estimated",
    "Father Name",
    "Region/State",
    "District",
    "Township (T)",
    "Village",
    "Ward",
    "Ward / Village tract",
    "Unique ID (UPI) - Index Case",
    "Relationship with index",
    "created",
    "lastUpdated",
    "enrollment_status",
    "program_id",
    "program_name",
    # 3. TB Screening Columns
    "[TB Screening] Age (at screening)",
    "[TB Screening] Any TB drug resistance history?",
    "[TB Screening] BMI",
    "[TB Screening] Breathlessness",
    "[TB Screening] CXR result",
    "[TB Screening] CXR result category",
    "[TB Screening] CXR screening date",
    "[TB Screening] CXR screening done",
    "[TB Screening] CXR screening facility type",
    "[TB Screening] Chest pain",
    "[TB Screening] Cough more than 2 weeks",
    "[TB Screening] Current activity",
    "[TB Screening] Enroll to Diagnostic Evaluation",
    "[TB Screening] Fatigue and Tiredness",
    "[TB Screening] Fever more than 2 weeks",
    "[TB Screening] Haemoptysis",
    "[TB Screening] Healthcare worker population",
    "[TB Screening] Height (in inches)",
    "[TB Screening] Household Contact",
    "[TB Screening] Loss of appetite",
    "[TB Screening] Migrant population",
    "[TB Screening] Night sweats",
    "[TB Screening] No symptoms related with TB",
    "[TB Screening] Number of previous TB episodes",
    "[TB Screening] Other referral organisation (Specify)",
    "[TB Screening] Other symptoms related with TB",
    "[TB Screening] Previous TB History",
    "[TB Screening] Previous TB regimen",
    "[TB Screening] Referral activity",
    "[TB Screening] Referral organization",
    "[TB Screening] Specify other symptoms",
    "[TB Screening] TB CS - HIV infection",
    "[TB Screening] TB CS - HIV status date",
    "[TB Screening] TB CS - Registration - Type of patient in last TB Treatment",
    "[TB Screening] TB CS - Risk factor alcohol",
    "[TB Screening] TB CS - Risk factor diabetes",
    "[TB Screening] TB CS - Risk factor smoking",
    "[TB Screening] TB CS - Risk factor undernourishment",
    "[TB Screening] Type of CXR",
    "[TB Screening] Weight (in kgs)",
    "[TB Screening] Weight loss",
    "[TB Screening] Year of last TB Treatment",
    # 4. TB Treatment & Outcome Columns
    "[2. TB Treatment] TB CS - Diagnosis date",
    "[2. TB Treatment] TB CS - First-line treatment regimen composition",
    "[2. TB Treatment] TB CS - First-line treatment start date",
    "[2. TB Treatment] TB CS - Manually assigned resistance classification",
    "[2. TB Treatment] TB CS - Outcome due date",
    "[2. TB Treatment] TB CS - Reassign resistance classification",
    "[2. TB Treatment] TB CS - Resistance at diagnosis",
    "[2. TB Treatment] TB CS - Resistance classification",
    "[2. TB Treatment] TB CS - Treatment initiation delay (days)",
    "[2. TB Treatment] TB CS - Treatment regimen",
    "[4. Outcome] TB CS - Treatment outcome",
    "[4. Outcome] TB CS - Treatment outcome delay (weeks)",
]

UID_REGEX = re.compile(r"^[A-Za-z][A-Za-z0-9]{10}$")
NULL_STRINGS = {"", "nan", "None", "NaN", "<NA>", "null"}

LEGACY_FIELDS = (
    "trackedEntityInstance,orgUnit,orgUnitName,created,lastUpdated,"
    "attributes[attribute,displayName,value],"
    "enrollments[enrollment,program,orgUnit,orgUnitName,enrolledAt,enrollmentDate,occurredAt,status,"
    "events[event,programStage,occurredAt,eventDate,created,status,dataValues[dataElement,value]]]"
)

MODERN_FIELDS = (
    "trackedEntity,orgUnit,orgUnitName,createdAt,updatedAt,"
    "attributes[attribute,displayName,value],"
    "enrollments[enrollment,program,orgUnit,orgUnitName,enrolledAt,occurredAt,status,"
    "events[event,programStage,occurredAt,createdAt,status,dataValues[dataElement,value]]]"
)


# --- Helper & Backend Functions ---
def normalize_base_url(url: str) -> str:
    """Ensures base_url is cleanly formatted without trailing slashes or duplicate /api paths."""
    clean_url = (url or "").strip().rstrip("/")
    if clean_url.endswith("/api"):
        clean_url = clean_url[:-4]
    return clean_url


def create_http_session(username: str, password: str, verify_ssl: bool = True) -> requests.Session:
    """Creates a pooled HTTP session with automatic retries and SSL fallback support."""
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({
        "Accept": "application/json",
        "User-Agent": "DHIS2-Streamlit-Exporter/8.0",
    })
    retries = Retry(
        total=3,
        connect=2,
        read=2,
        backoff_factor=0.5,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=frozenset(["GET"]),
    )
    adapter = HTTPAdapter(pool_connections=16, pool_maxsize=16, max_retries=retries)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def safe_get(session: requests.Session, url: str, **kwargs) -> requests.Response:
    """Performs GET request and automatically falls back to verify=False on SSL errors."""
    try:
        return session.get(url, **kwargs)
    except requests.exceptions.SSLError:
        session.verify = False
        return session.get(url, verify=False, **kwargs)


def reorder_columns(df: pd.DataFrame) -> pd.DataFrame:
    """
    Orders any DataFrame so that Date, OrgUnit, UPI, Name, Age, Sex appear first,
    followed by all remaining columns (excluding internal helper columns).
    """
    if df.empty and len(df.columns) == 0:
        return df
    internal_cols = {
        "enrollment_id",
        "tei_orgUnit_id",
        "_upi_sort",
        "_date_sort",
        "_prog_sort",
        "_created_sort",
        "_effective_key",
    }
    # Deduplicate column names if any exist while preserving order
    if df.columns.duplicated().any():
        df = df.loc[:, ~df.columns.duplicated(keep="first")]

    leading = [col for col in PRIORITY_LEADING_COLUMNS if col in df.columns]
    remaining = [col for col in df.columns if col not in leading and col not in internal_cols]
    return df[leading + remaining]


def sort_patient_records_chronologically(df: pd.DataFrame, join_key: str) -> pd.DataFrame:
    """Sorts records by Unique ID (UPI), enrollment_date, and clinical cascade order."""
    if df.empty or join_key not in df.columns:
        return df

    sorted_df = df.copy()
    sorted_df["_upi_sort"] = (
        sorted_df[join_key]
        .fillna("~~~~")
        .astype(str)
        .str.strip()
        .replace({"": "~~~~", "None": "~~~~", "nan": "~~~~"})
    )

    if "enrollment_date" in sorted_df.columns:
        first_date_str = (
            sorted_df["enrollment_date"]
            .fillna("")
            .astype(str)
            .str.split("|")
            .str[0]
            .str.strip()
            .str.slice(0, 10)
        )
        sorted_df["_date_sort"] = pd.to_datetime(first_date_str, format="%Y-%m-%d", errors="coerce")
    else:
        sorted_df["_date_sort"] = pd.NaT

    if "program_id" in sorted_df.columns:
        first_prog = sorted_df["program_id"].fillna("").astype(str).str.split("|").str[0].str.strip()
        sorted_df["_prog_sort"] = first_prog.map(PROGRAM_CASCADE_ORDER).fillna(99)
    elif "program_name" in sorted_df.columns:
        first_prog_name = sorted_df["program_name"].fillna("").astype(str).str.split("|").str[0].str.strip()
        sorted_df["_prog_sort"] = first_prog_name.map(PROGRAM_CASCADE_ORDER).fillna(99)
    else:
        sorted_df["_prog_sort"] = 99

    if "created" in sorted_df.columns:
        sorted_df["_created_sort"] = sorted_df["created"].fillna("").astype(str)
    else:
        sorted_df["_created_sort"] = ""

    sorted_df = (
        sorted_df.sort_values(
            by=["_upi_sort", "_date_sort", "_prog_sort", "_created_sort"],
            ascending=[True, True, True, True],
            na_position="last",
            kind="mergesort",
        )
        .drop(columns=["_upi_sort", "_date_sort", "_prog_sort", "_created_sort"], errors="ignore")
        .reset_index(drop=True)
    )
    return sorted_df


def clean_orgunit_name_series(series: pd.Series, meta_map: dict[str, str] | None = None) -> pd.Series:
    """Uncodes raw 11-character DHIS2 UIDs in `Org Unit Name` into display values."""
    if series.empty:
        return series

    def _clean_val(val: Any) -> str | None:
        if pd.isna(val):
            return None
        val_s = str(val).strip()
        if val_s in NULL_STRINGS:
            return None
        parts = [p.strip() for p in val_s.split("|") if p.strip() and p.strip() not in NULL_STRINGS]
        if not parts:
            return None
        resolved = [meta_map.get(p, p) if meta_map else p for p in parts]
        unique_parts = list(dict.fromkeys(resolved))
        named_parts = [p for p in unique_parts if not UID_REGEX.match(p)]
        if named_parts:
            return " | ".join(named_parts)
        return " | ".join(unique_parts)

    return series.map(_clean_val)


@st.cache_data(ttl=3600, show_spinner=False, max_entries=1)
def fetch_cached_metadata(
    base_url: str, username: str, password: str
) -> tuple[dict[str, str], dict[str, dict[str, str]], list[str], bool, bool]:
    """
    Fetches and caches:
    1. `meta_map`: UID/Code -> Display Name for Programs, Stages, DataElements, Attributes, and OrgUnits.
    2. `option_value_map`: Field-specific OptionSet code -> Display Name mapping for uncoding coded values.
    3. `user_ous`: Accessible Organization Units for the user account/role.
    4. `is_legacy`: Whether the server uses the legacy or modern Tracker API.
    5. `verify_ssl`: Whether SSL verification succeeded or required fallback.

    Raises RuntimeError on authentication or connection failure so failed calls are NEVER cached.
    """
    session = create_http_session(username, password, verify_ssl=True)
    meta_map: dict[str, str] = {}
    option_value_map: dict[str, dict[str, str]] = {}

    # Step 0: Verify authentication & fetch user OrgUnits via /api/me.json first
    me_url = (
        f"{base_url}/api/me.json?"
        "fields=id,username,organisationUnits[id,displayName],"
        "dataViewOrganisationUnits[id,displayName],"
        "teiSearchOrganisationUnits[id,displayName]"
    )
    try:
        me_res = safe_get(session, me_url, timeout=15)
    except requests.exceptions.RequestException as exc:
        session.close()
        raise RuntimeError(f"Unable to connect to DHIS2 server at {base_url}: {exc}") from exc

    if me_res.status_code in (401, 403):
        session.close()
        raise RuntimeError(
            f"Authentication failed (HTTP {me_res.status_code}). Please check your DHIS2 username and password."
        )
    if not me_res.ok:
        session.close()
        raise RuntimeError(
            f"DHIS2 server returned HTTP {me_res.status_code} when checking `/api/me.json`."
        )

    verify_ssl = bool(session.verify)
    user_ous: set[str] = set()
    try:
        me_data = me_res.json()
        for key in ("teiSearchOrganisationUnits", "dataViewOrganisationUnits", "organisationUnits"):
            for item in me_data.get(key, []):
                uid = item.get("id")
                if uid:
                    user_ous.add(uid)
                    if item.get("displayName"):
                        meta_map[uid] = item["displayName"]
    except Exception:
        pass

    # Pre-seed configured Township OUs so they always resolve even if full orgUnit list is huge
    for t_name, uids in TOWNSHIP_OU_MAP.items():
        for uid in uids:
            meta_map[uid] = t_name

    endpoints = {
        "trackedEntityAttributes": (
            f"{base_url}/api/trackedEntityAttributes.json?"
            "fields=id,displayName,optionSet[id,options[code,name,displayName]]&paging=false"
        ),
        "dataElements": (
            f"{base_url}/api/dataElements.json?"
            "fields=id,displayName,optionSet[id,options[code,name,displayName]]&paging=false"
        ),
        "programStages": f"{base_url}/api/programStages.json?fields=id,displayName&paging=false",
        "programs": f"{base_url}/api/programs.json?fields=id,displayName&paging=false",
        "organisationUnits": f"{base_url}/api/organisationUnits.json?fields=id,code,displayName&paging=false",
    }

    def _fetch_meta(resource_url: tuple[str, str]) -> tuple[str, list[dict]]:
        resource, url = resource_url
        try:
            res = safe_get(session, url, timeout=30)
            if res.ok:
                return resource, res.json().get(resource, [])
        except Exception:
            pass
        return resource, []

    with ThreadPoolExecutor(max_workers=3) as executor:
        for resource, items in executor.map(_fetch_meta, endpoints.items()):
            for item in items:
                uid = item.get("id")
                if not uid:
                    continue
                display_name = item.get("displayName") or item.get("code") or uid
                meta_map[uid] = display_name

                if resource == "organisationUnits" and item.get("code"):
                    meta_map[str(item["code"]).strip()] = display_name

                opt_set = item.get("optionSet")
                if opt_set and isinstance(opt_set, dict):
                    field_opts: dict[str, str] = {}
                    for opt in opt_set.get("options", []) or []:
                        if not isinstance(opt, dict):
                            continue
                        opt_code = str(opt.get("code", "")).strip()
                        opt_label = opt.get("displayName") or opt.get("name") or opt_code
                        if opt_code:
                            field_opts[opt_code] = opt_label
                    if field_opts:
                        option_value_map[uid] = field_opts

    # Re-apply Township OU labels if any remained raw UIDs
    for t_name, uids in TOWNSHIP_OU_MAP.items():
        for uid in uids:
            if uid not in meta_map or UID_REGEX.match(meta_map[uid]):
                meta_map[uid] = t_name

    # Probe whether legacy `/api/trackedEntityInstances.json` is supported (include program param!)
    is_legacy = True
    try:
        test_res = safe_get(
            session,
            f"{base_url}/api/trackedEntityInstances.json",
            params={
                "program": REGISTRATION_PROGRAM_ID,
                "ouMode": "ACCESSIBLE",
                "pageSize": 1,
            },
            timeout=12,
        )
        if test_res.status_code in (404, 410):
            is_legacy = False
    except Exception:
        pass

    session.close()
    return meta_map, option_value_map, list(user_ous), is_legacy, verify_ssl


def _build_fetch_params(
    is_legacy: bool,
    prog: str,
    ou: str | None,
    ou_mode: str,
    page: int,
    page_size: int,
    use_modern_org_unit_mode: bool = False,
) -> dict[str, Any]:
    """Builds DHIS2 API query parameters for legacy or modern tracker endpoints."""
    if is_legacy:
        params: dict[str, Any] = {
            "program": prog,
            "ouMode": ou_mode,
            "pageSize": page_size,
            "page": page,
            "totalPages": "false",
            "fields": LEGACY_FIELDS,
        }
        if ou:
            params["ou"] = ou
    else:
        mode_key = "orgUnitMode" if use_modern_org_unit_mode else "ouMode"
        params = {
            "program": prog,
            "pageSize": page_size,
            "page": page,
            "totalPages": "false",
            "fields": MODERN_FIELDS,
            mode_key: ou_mode,
        }
        if ou:
            params["orgUnits" if use_modern_org_unit_mode else "orgUnit"] = ou
    return params


def _fetch_single_page(
    session: requests.Session,
    active_endpoint: str,
    is_legacy: bool,
    prog: str,
    ou: str | None,
    ou_mode: str,
    page: int,
    page_size: int = 500,
) -> list[dict]:
    """Fetches a single page of tracked entities and tags each with the queried program."""
    params = _build_fetch_params(is_legacy, prog, ou, ou_mode, page, page_size, use_modern_org_unit_mode=False)
    try:
        response = safe_get(session, active_endpoint, params=params, timeout=35)
        if not response.ok and not is_legacy:
            alt_params = _build_fetch_params(
                is_legacy, prog, ou, ou_mode, page, page_size, use_modern_org_unit_mode=True
            )
            response = safe_get(session, active_endpoint, params=alt_params, timeout=35)
        if not response.ok:
            return []
        data = response.json()
        instances = (
            data.get("trackedEntityInstances")
            or data.get("instances")
            or data.get("trackedEntities")
            or []
        )
        for inst in instances:
            inst["_queried_program"] = prog
        return instances
    except Exception:
        return []


def fetch_raw_program_pages(
    session: requests.Session,
    active_endpoint: str,
    is_legacy: bool,
    prog: str,
    ou: str | None = None,
    ou_mode: str = "ACCESSIBLE",
) -> list[dict]:
    """Fetches all Tracked Entity pages in bounded parallel batches."""
    page_size = 500

    first_page = _fetch_single_page(session, active_endpoint, is_legacy, prog, ou, ou_mode, 1, page_size)
    if not first_page:
        return []
    if len(first_page) < page_size:
        return first_page

    all_instances = list(first_page)
    _fetch_page = partial(
        _fetch_single_page, session, active_endpoint, is_legacy, prog, ou, ou_mode, page_size=page_size
    )

    batch_size = 3
    page = 2
    while True:
        pages_in_batch = list(range(page, page + batch_size))
        with ThreadPoolExecutor(max_workers=batch_size) as page_executor:
            batch_results = list(page_executor.map(_fetch_page, pages_in_batch))

        got_last_page = False
        for result in batch_results:
            if not result:
                got_last_page = True
                break
            all_instances.extend(result)
            if len(result) < page_size:
                got_last_page = True
                break

        if got_last_page:
            break
        page += batch_size

    return all_instances


def build_uncoded_df_dhis2(
    data_dhis: list[dict],
    meta_map: dict[str, str],
    option_value_map: dict[str, dict[str, str]],
) -> tuple[pd.DataFrame, dict[str, str]]:
    """
    Transforms raw `data_dhis` records into the uncoded source DataFrame `df_dhis2`
    by replacing metadata UIDs and option codes with display values.
    Does not mutate the cached `meta_map` in place.
    """
    local_meta = dict(meta_map)
    uncoded_rows: list[dict[str, Any]] = []
    _prog_label = PROGRAM_ID_TO_LABEL
    _meta_get = local_meta.get
    _opt_map = option_value_map
    _uid_match = UID_REGEX.match

    for instance in data_dhis:
        prog = instance.get("_queried_program") or ""
        prog_name = _prog_label.get(prog) or _meta_get(prog, prog)
        tei_id = instance.get("trackedEntityInstance") or instance.get("trackedEntity")
        tei_ou_id = instance.get("orgUnit")
        if tei_ou_id and (ou_name_raw := instance.get("orgUnitName")):
            local_meta[tei_ou_id] = ou_name_raw
        tei_ou_name = instance.get("orgUnitName") or _meta_get(tei_ou_id, tei_ou_id)

        # Uncode Tracked Entity Attributes
        base_attrs: dict[str, Any] = {}
        for attr in instance.get("attributes", []) or []:
            attr_id = attr.get("attribute") or ""
            val = attr.get("value")
            if val is None:
                continue
            val_str = str(val).strip()
            if not val_str or val_str in NULL_STRINGS:
                continue
            if attr_id in _opt_map:
                decoded_val = _opt_map[attr_id].get(val_str, val_str)
            elif _uid_match(val_str) and val_str in local_meta:
                decoded_val = local_meta[val_str]
            else:
                decoded_val = val_str
            col_name = attr.get("displayName") or _meta_get(attr_id) or attr_id
            base_attrs[col_name] = decoded_val

        enrollments = [
            enr for enr in (instance.get("enrollments") or [])
            if enr.get("program") == prog
        ]

        inst_created = instance.get("created") or instance.get("createdAt")
        inst_updated = instance.get("lastUpdated") or instance.get("updatedAt")

        if not enrollments:
            row = {
                **base_attrs,
                "program_id": prog,
                "program_name": prog_name,
                "trackedEntityInstance": tei_id,
                "enrollment_id": f"no_enr_{tei_id}",
                "orgUnit_id": tei_ou_id,
                "tei_orgUnit_id": tei_ou_id,
                "Org Unit Name": _meta_get(tei_ou_id, tei_ou_name),
                "created": inst_created,
                "lastUpdated": inst_updated,
            }
            uncoded_rows.append(row)
        else:
            for enrollment in enrollments:
                enr_id = enrollment.get("enrollment")
                enr_ou_id = enrollment.get("orgUnit") or tei_ou_id
                if enr_ou_id and (enr_ou_name_raw := enrollment.get("orgUnitName")):
                    local_meta[enr_ou_id] = enr_ou_name_raw

                enr_ou_name = (
                    enrollment.get("orgUnitName")
                    or _meta_get(enr_ou_id)
                    or _meta_get(tei_ou_id)
                    or enr_ou_id
                )
                raw_enr_date = (
                    enrollment.get("enrolledAt")
                    or enrollment.get("enrollmentDate")
                    or enrollment.get("occurredAt")
                )
                enr_date_clean = str(raw_enr_date)[:10] if raw_enr_date else None

                # System keys placed AFTER **base_attrs so attributes never overwrite system metadata
                row = {
                    **base_attrs,
                    "program_id": prog,
                    "program_name": prog_name,
                    "trackedEntityInstance": tei_id,
                    "enrollment_id": enr_id,
                    "orgUnit_id": enr_ou_id,
                    "tei_orgUnit_id": tei_ou_id,
                    "Org Unit Name": enr_ou_name,
                    "created": inst_created,
                    "lastUpdated": inst_updated,
                    "enrollment_date": enr_date_clean,
                    "enrollment_status": enrollment.get("status"),
                }

                # Sort events chronologically before aggregating data values
                events = enrollment.get("events") or []
                if len(events) > 1:
                    events = sorted(
                        events,
                        key=lambda ev: str(
                            ev.get("occurredAt")
                            or ev.get("eventDate")
                            or ev.get("created")
                            or ev.get("createdAt")
                            or ""
                        ),
                    )

                for event in events:
                    stage_id = event.get("programStage") or ""
                    stage_name = _meta_get(stage_id, f"Stage_{stage_id}")

                    for dv in event.get("dataValues") or []:
                        de_id = dv.get("dataElement") or ""
                        val = dv.get("value")
                        if val is None:
                            continue
                        val_str = str(val).strip()
                        if not val_str or val_str in NULL_STRINGS:
                            continue
                        if de_id in _opt_map:
                            decoded_val = _opt_map[de_id].get(val_str, val_str)
                        elif _uid_match(val_str) and val_str in local_meta:
                            decoded_val = local_meta[val_str]
                        else:
                            decoded_val = val_str

                        de_name = _meta_get(de_id, f"Element_{de_id}")
                        col_key = f"[{stage_name}] {de_name}"

                        existing = row.get(col_key)
                        if existing:
                            if decoded_val not in existing.split(" | "):
                                row[col_key] = f"{existing} | {decoded_val}"
                        else:
                            row[col_key] = decoded_val

                uncoded_rows.append(row)

    df_dhis2 = pd.DataFrame(uncoded_rows)
    if not df_dhis2.empty:
        if "orgUnit_id" in df_dhis2.columns and "Org Unit Name" in df_dhis2.columns:
            df_dhis2["Org Unit Name"] = (
                df_dhis2["orgUnit_id"].map(local_meta).fillna(df_dhis2["Org Unit Name"])
            )
            df_dhis2["Org Unit Name"] = clean_orgunit_name_series(
                df_dhis2["Org Unit Name"], local_meta
            )

        dedup_cols = [
            c for c in ["program_id", "trackedEntityInstance", "enrollment_id"]
            if c in df_dhis2.columns
        ]
        if dedup_cols:
            df_dhis2 = df_dhis2.drop_duplicates(subset=dedup_cols, keep="first").reset_index(drop=True)

    return df_dhis2, local_meta


@st.cache_data(ttl=1800, show_spinner=False, max_entries=1)
def get_cached_source_df_dhis2(
    web: str,
    username: str,
    password: str,
    prog_ids_tuple: tuple[str, ...],
) -> tuple[pd.DataFrame, dict[str, str], int]:
    """
    Stage 1 & 2 (Cached):
    1. Extracts all raw data from DHIS2 (`data_dhis`).
    2. Creates and caches the uncoded source DataFrame (`df_dhis2`).
    """
    base_url = normalize_base_url(web)
    meta_map, option_value_map, user_ous, is_legacy, verify_ssl = fetch_cached_metadata(
        base_url, username, password
    )
    active_endpoint = (
        f"{base_url}/api/trackedEntityInstances.json"
        if is_legacy
        else f"{base_url}/api/tracker/trackedEntities.json"
    )

    all_configured_ous = list({ou for ous in TOWNSHIP_OU_MAP.values() for ou in ous})
    candidate_ous = list(dict.fromkeys(all_configured_ous + user_ous))

    session = create_http_session(username, password, verify_ssl=verify_ssl)
    data_dhis: list[dict] = []

    def _fetch_program(prog: str) -> list[dict]:
        prog_instances = fetch_raw_program_pages(
            session=session,
            active_endpoint=active_endpoint,
            is_legacy=is_legacy,
            prog=prog,
            ou=None,
            ou_mode="ACCESSIBLE",
        )
        if not prog_instances and candidate_ous:
            # Sequential OU fallback avoids nested ThreadPoolExecutor explosion
            for ou in candidate_ous:
                ou_instances = fetch_raw_program_pages(
                    session=session,
                    active_endpoint=active_endpoint,
                    is_legacy=is_legacy,
                    prog=prog,
                    ou=ou,
                    ou_mode="DESCENDANTS",
                )
                if ou_instances:
                    prog_instances.extend(ou_instances)
        return prog_instances

    max_prog_workers = max(1, min(2, len(prog_ids_tuple)))
    with ThreadPoolExecutor(max_workers=max_prog_workers) as prog_executor:
        futures = [prog_executor.submit(_fetch_program, p) for p in prog_ids_tuple]
        for future in as_completed(futures):
            data_dhis.extend(future.result())

    session.close()

    raw_tei_count = len(data_dhis)
    df_dhis2, updated_meta_map = build_uncoded_df_dhis2(data_dhis, meta_map, option_value_map)
    del data_dhis
    gc.collect()
    return df_dhis2, updated_meta_map, raw_tei_count


def fast_combine_by_key(
    df: pd.DataFrame,
    join_key: str,
    column_order: list[str],
    meta_map: dict[str, str] | None = None,
) -> pd.DataFrame:
    """Vectorized deduplication and pipe-concatenation by `join_key`."""
    if df.empty or join_key not in df.columns:
        return pd.DataFrame(columns=column_order)

    # Deduplicate column names if needed
    if df.columns.duplicated().any():
        df = df.loc[:, ~df.columns.duplicated(keep="first")]

    valid_mask = df[join_key].notna() & ~df[join_key].astype(str).str.strip().isin(NULL_STRINGS)
    valid_df = df[valid_mask].copy()
    if valid_df.empty:
        return pd.DataFrame(columns=column_order)

    valid_df = sort_patient_records_chronologically(valid_df, join_key)

    dup_mask = valid_df.duplicated(subset=[join_key], keep=False)
    if not dup_mask.any():
        out = valid_df.reindex(columns=column_order).reset_index(drop=True)
        if "Org Unit Name" in out.columns:
            out["Org Unit Name"] = clean_orgunit_name_series(out["Org Unit Name"], meta_map)
        return out

    df_unique = valid_df[~dup_mask].reindex(columns=column_order)
    df_dups = valid_df[dup_mask]

    # Preserve all duplicate keys so no patient is ever lost even if all non-key columns are empty
    base_dup_keys = df_dups[[join_key]].drop_duplicates(keep="first").reset_index(drop=True)
    non_empty_cols = [c for c in df_dups.columns if c != join_key and df_dups[c].notna().any()]

    if not non_empty_cols:
        df_dups_agg = base_dup_keys.reindex(columns=column_order)
    else:
        melted = (
            df_dups[[join_key] + non_empty_cols]
            .melt(id_vars=join_key, var_name="_col", value_name="_val")
            .dropna(subset=["_val"])
        )
        melted["_val"] = melted["_val"].astype(str).str.strip()
        melted = melted[~melted["_val"].isin(NULL_STRINGS)]

        if melted.empty:
            df_dups_agg = base_dup_keys.reindex(columns=column_order)
        else:
            if melted["_val"].str.contains("|", regex=False).any():
                melted["_val"] = melted["_val"].str.split(r"\s*\|\s*")
                melted = melted.explode("_val")
                melted["_val"] = melted["_val"].astype(str).str.strip()
                melted = melted[~melted["_val"].isin(NULL_STRINGS)]

            melted = melted.drop_duplicates(subset=[join_key, "_col", "_val"], keep="first")

            if melted.empty:
                df_dups_agg = base_dup_keys.reindex(columns=column_order)
            else:
                pivoted = (
                    melted.groupby([join_key, "_col"], sort=False)["_val"]
                    .agg(" | ".join)
                    .unstack("_col")
                    .reset_index()
                )
                df_dups_agg = (
                    base_dup_keys.merge(pivoted, on=join_key, how="left")
                    .reindex(columns=column_order)
                )

    combined = pd.concat([df_unique, df_dups_agg], ignore_index=True)
    if "Org Unit Name" in combined.columns:
        combined["Org Unit Name"] = clean_orgunit_name_series(combined["Org Unit Name"], meta_map)

    return sort_patient_records_chronologically(combined, join_key)


def count_unique_upis(df: pd.DataFrame, join_key: str) -> int:
    if df.empty or join_key not in df.columns:
        return 0
    s = df[join_key].dropna().astype(str).str.strip()
    s = s[~s.isin(NULL_STRINGS)]
    return int(s.nunique())


def sanitize_df_for_display(df: pd.DataFrame) -> pd.DataFrame:
    """
    Ensures all object/mixed columns in a DataFrame are clean strings or None
    so Streamlit / PyArrow never crashes with ArrowTypeError on mixed types.
    """
    if df.empty:
        return df
    clean = df.copy()
    for col in clean.columns:
        if pd.api.types.is_object_dtype(clean[col]):
            clean[col] = clean[col].where(clean[col].notna(), None)
            clean[col] = clean[col].map(lambda x: str(x) if x is not None else None)
    return clean


def split_and_process_dataframes(
    df_dhis2: pd.DataFrame,
    selected_ou_ids: list[str],
    selected_township_names: list[str] | None = None,
    start_date_str: str | None = None,
    end_date_str: str | None = None,
    meta_map: dict[str, str] | None = None,
) -> dict[str, Any]:
    """
    Stage 3, 4 & 5:
    1. Splits `df_dhis2` into page DataFrames (`df_registration`, `df_diagnosis`, `df_surveillance`, `df_contact`).
    2. Filters `df_registration` by selected OrgUnits & dates and deduplicates by UPI.
    3. Filters `df_diagnosis`, `df_surveillance`, `df_contact` by all unique UPIs of `df_registration`
       regardless of OrgUnit, keeping multiple records for different episodes.
    4. Creates `df_combined` (accumulated multi-episode events) and `df_ygntbpro` (merged by UPI & column-filtered).
    """
    if df_dhis2.empty:
        empty_df = pd.DataFrame(columns=PRIORITY_LEADING_COLUMNS)
        empty_summary = pd.DataFrame(
            columns=["DataFrame / Page", "Raw Server Rows", "Processed Cohort Rows", "Unique UPIs", "Logic Applied"]
        )
        return {
            "df_registration": empty_df,
            "df_diagnosis": empty_df,
            "df_dignosis": empty_df,
            "df_surveillance": empty_df,
            "df_contact": empty_df,
            "df_combined": empty_df,
            "df_ygntbpro": empty_df,
            "df_summary": empty_summary,
        }

    df_dhis2_clean = df_dhis2.copy()
    if df_dhis2_clean.columns.duplicated().any():
        df_dhis2_clean = df_dhis2_clean.loc[:, ~df_dhis2_clean.columns.duplicated(keep="first")]

    join_key = "Unique ID (UPI)" if "Unique ID (UPI)" in df_dhis2_clean.columns else "trackedEntityInstance"
    if join_key not in df_dhis2_clean.columns:
        df_dhis2_clean[join_key] = None

    df_dhis2_clean[join_key] = df_dhis2_clean[join_key].astype(str).str.strip()
    df_dhis2_clean.loc[df_dhis2_clean[join_key].isin(NULL_STRINGS), join_key] = None

    all_cols = list(df_dhis2_clean.columns)

    # --- Step 1: Split `df_dhis2` by Pages ---
    if "program_id" in df_dhis2_clean.columns:
        df_registration_raw = df_dhis2_clean[df_dhis2_clean["program_id"] == REGISTRATION_PROGRAM_ID].copy()
        df_diagnosis_raw = df_dhis2_clean[df_dhis2_clean["program_id"] == DIAGNOSIS_PROGRAM_ID].copy()
        df_surveillance_raw = df_dhis2_clean[df_dhis2_clean["program_id"] == SURVEILLANCE_PROGRAM_ID].copy()
        df_contact_raw = df_dhis2_clean[df_dhis2_clean["program_id"] == CONTACT_PROGRAM_ID].copy()
    else:
        df_registration_raw = pd.DataFrame(columns=all_cols)
        df_diagnosis_raw = pd.DataFrame(columns=all_cols)
        df_surveillance_raw = pd.DataFrame(columns=all_cols)
        df_contact_raw = pd.DataFrame(columns=all_cols)

    # --- Step 2: Process `df_registration` (Filter Selected OrgUnits + Unique UPI) ---
    df_reg_filtered = df_registration_raw.copy()
    if selected_ou_ids and not df_reg_filtered.empty:
        ou_mask = pd.Series(False, index=df_reg_filtered.index)
        if "orgUnit_id" in df_reg_filtered.columns:
            ou_mask |= df_reg_filtered["orgUnit_id"].isin(selected_ou_ids)
        if "tei_orgUnit_id" in df_reg_filtered.columns:
            ou_mask |= df_reg_filtered["tei_orgUnit_id"].isin(selected_ou_ids)
        if selected_township_names and "Org Unit Name" in df_reg_filtered.columns:
            ou_mask |= df_reg_filtered["Org Unit Name"].isin(selected_township_names)
        df_reg_filtered = df_reg_filtered[ou_mask].copy()

    if (start_date_str or end_date_str) and "enrollment_date" in df_reg_filtered.columns and not df_reg_filtered.empty:
        enr_dates = df_reg_filtered["enrollment_date"].fillna("").astype(str).str.slice(0, 10)
        if start_date_str:
            df_reg_filtered = df_reg_filtered[enr_dates >= start_date_str]
            enr_dates = df_reg_filtered["enrollment_date"].fillna("").astype(str).str.slice(0, 10)
        if end_date_str:
            df_reg_filtered = df_reg_filtered[enr_dates <= end_date_str]

    df_registration = fast_combine_by_key(df_reg_filtered, join_key, all_cols, meta_map)
    registration_upis = (
        set(df_registration[join_key].dropna().unique())
        if join_key in df_registration.columns and not df_registration.empty
        else set()
    )

    # --- Step 3: Process `df_diagnosis`, `df_surveillance`, `df_contact` ---
    if registration_upis and join_key in df_dhis2_clean.columns:
        df_diagnosis = sort_patient_records_chronologically(
            df_diagnosis_raw[df_diagnosis_raw[join_key].isin(registration_upis)].reindex(columns=all_cols),
            join_key,
        )
        df_surveillance = sort_patient_records_chronologically(
            df_surveillance_raw[df_surveillance_raw[join_key].isin(registration_upis)].reindex(columns=all_cols),
            join_key,
        )
        df_contact = sort_patient_records_chronologically(
            df_contact_raw[df_contact_raw[join_key].isin(registration_upis)].reindex(columns=all_cols),
            join_key,
        )
    else:
        df_diagnosis = pd.DataFrame(columns=all_cols)
        df_surveillance = pd.DataFrame(columns=all_cols)
        df_contact = pd.DataFrame(columns=all_cols)

    # --- Step 4: Create `df_combined` (Accumulated Multi-Episode Events Across Pages) ---
    df_combined = pd.concat(
        [df_registration, df_diagnosis, df_surveillance, df_contact],
        ignore_index=True,
    )
    if "Org Unit Name" in df_combined.columns:
        df_combined["Org Unit Name"] = clean_orgunit_name_series(df_combined["Org Unit Name"], meta_map)
    df_combined = sort_patient_records_chronologically(df_combined, join_key)

    # Apply priority column ordering (Date, OrgUnit, UPI, Name, Age, Sex -> Remaining)
    df_registration = reorder_columns(df_registration)
    df_diagnosis = reorder_columns(df_diagnosis)
    df_surveillance = reorder_columns(df_surveillance)
    df_contact = reorder_columns(df_contact)
    df_combined = reorder_columns(df_combined)

    # --- Step 5: Create `df_ygntbpro` (Filter Columns First, Then Merge by UPI for Speed) ---
    # Ensure sorting helper columns are available during combine even if not in SELECTED_COLUMN_LIST
    helper_sort_cols = [
        c for c in ["enrollment_date", "program_id", "program_name", "created"]
        if c in df_combined.columns
    ]
    candidate_ygn_cols = list(
        dict.fromkeys(
            [join_key]
            + [col for col in SELECTED_COLUMN_LIST if col in df_combined.columns]
            + helper_sort_cols
        )
    )
    df_combined_slim = df_combined[candidate_ygn_cols].copy() if candidate_ygn_cols else df_combined.copy()
    df_ygntbpro = fast_combine_by_key(df_combined_slim, join_key, candidate_ygn_cols, meta_map)

    # Keep only the requested YgnTBPro columns (plus join_key if fallback was used)
    final_ygn_cols = list(
        dict.fromkeys(
            [col for col in SELECTED_COLUMN_LIST if col in df_ygntbpro.columns]
            + ([join_key] if join_key not in SELECTED_COLUMN_LIST and join_key in df_ygntbpro.columns else [])
        )
    )
    df_ygntbpro = df_ygntbpro.reindex(columns=final_ygn_cols)
    df_ygntbpro = reorder_columns(df_ygntbpro)
    df_ygntbpro = sort_patient_records_chronologically(df_ygntbpro, join_key)

    # --- Step 6: Build Summary Table ---
    total_source_rows = len(df_dhis2_clean)
    total_source_upis = count_unique_upis(df_dhis2_clean, join_key)
    summary_rows = [
        {
            "DataFrame / Page": "df_dhis2 (Uncoded Source Cache)",
            "Raw Server Rows": total_source_rows,
            "Processed Cohort Rows": total_source_rows,
            "Unique UPIs": total_source_upis,
            "Logic Applied": "Uncoded Display Values (Optional CSV Export)",
        },
        {
            "DataFrame / Page": "1. df_registration (Registration & Screening)",
            "Raw Server Rows": len(df_registration_raw),
            "Processed Cohort Rows": len(df_registration),
            "Unique UPIs": count_unique_upis(df_registration, join_key),
            "Logic Applied": "Filtered by Selected OrgUnits & Deduplicated by UPI",
        },
        {
            "DataFrame / Page": "2. df_diagnosis (Diagnostic Evaluation)",
            "Raw Server Rows": len(df_diagnosis_raw),
            "Processed Cohort Rows": len(df_diagnosis),
            "Unique UPIs": count_unique_upis(df_diagnosis, join_key),
            "Logic Applied": "Filtered by df_registration UPIs (All OUs, Multi-Episode Kept)",
        },
        {
            "DataFrame / Page": "3. df_surveillance (TB Surveillance)",
            "Raw Server Rows": len(df_surveillance_raw),
            "Processed Cohort Rows": len(df_surveillance),
            "Unique UPIs": count_unique_upis(df_surveillance, join_key),
            "Logic Applied": "Filtered by df_registration UPIs (All OUs, Multi-Episode Kept)",
        },
        {
            "DataFrame / Page": "4. df_contact (TB Contact Tracing & TPT)",
            "Raw Server Rows": len(df_contact_raw),
            "Processed Cohort Rows": len(df_contact),
            "Unique UPIs": count_unique_upis(df_contact, join_key),
            "Logic Applied": "Filtered by df_registration UPIs (All OUs, Multi-Episode Kept)",
        },
        {
            "DataFrame / Page": "5. df_combined (Combined Multi-Event Sheet)",
            "Raw Server Rows": total_source_rows,
            "Processed Cohort Rows": len(df_combined),
            "Unique UPIs": count_unique_upis(df_combined, join_key),
            "Logic Applied": "Accumulated Pages (df_registration + Remaining Pages, Sorted)",
        },
        {
            "DataFrame / Page": "6. df_ygntbpro (YgnTBPro Summary Sheet)",
            "Raw Server Rows": total_source_rows,
            "Processed Cohort Rows": len(df_ygntbpro),
            "Unique UPIs": count_unique_upis(df_ygntbpro, join_key),
            "Logic Applied": f"Merged by UPI & Filtered ({len(df_ygntbpro.columns)} Selected Columns)",
        },
    ]
    df_summary = pd.DataFrame(summary_rows)

    del df_dhis2_clean, df_registration_raw, df_diagnosis_raw, df_surveillance_raw, df_contact_raw
    gc.collect()

    return {
        "df_registration": sanitize_df_for_display(df_registration),
        "df_diagnosis": sanitize_df_for_display(df_diagnosis),
        "df_dignosis": sanitize_df_for_display(df_diagnosis),
        "df_surveillance": sanitize_df_for_display(df_surveillance),
        "df_contact": sanitize_df_for_display(df_contact),
        "df_combined": sanitize_df_for_display(df_combined),
        "df_ygntbpro": sanitize_df_for_display(df_ygntbpro),
        "df_summary": df_summary,
    }


def build_excel_sheets_dict(res: dict[str, Any]) -> dict[str, pd.DataFrame]:
    """Builds the ordered sheet dictionary on demand without duplicating references in session_state."""
    return {
        "Summary": res["df_summary"],
        "YgnTBPro": res["df_ygntbpro"],
        "Combined": res["df_combined"],
        "Registration & Screening": res["df_registration"],
        "Diagnostic Evaluation": res["df_diagnosis"],
        "TB Surveillance": res["df_surveillance"],
        "TB Contact Tracing and TPT": res["df_contact"],
    }


def _sanitize_excel_value(val: Any) -> Any:
    """Strips illegal XML control characters that crash openpyxl and normalizes nulls."""
    if val is None or pd.isna(val):
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = ILLEGAL_CHARACTERS_RE.sub("", str(val))
    # Prevent Excel formula injection / accidental formula parse errors on strings starting with '='
    if s.startswith("="):
        s = "'" + s
    return s


def convert_df_to_excel_bytes(sheets_dict: dict[str, pd.DataFrame]) -> bytes:
    """Fast streaming Excel builder (`write_only=True`) with XML control-char sanitization."""
    output = io.BytesIO()
    wb = openpyxl.Workbook(write_only=True)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheets_written = 0
    for sheet_name, sheet_df in sheets_dict.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        sheets_written += 1

        if sheet_name != "Summary":
            sheet_df = reorder_columns(sheet_df)

        non_empty_cols = [
            c for c in sheet_df.columns
            if c in PRIORITY_LEADING_COLUMNS or sheet_df[c].notna().any()
        ]
        sub_df = sheet_df[non_empty_cols] if non_empty_cols else sheet_df
        headers = [str(h) for h in sub_df.columns]
        if not headers:
            ws.append(["No data available"])
            continue

        for col_idx, col_name in enumerate(headers, start=1):
            sample_series = sub_df.iloc[:300, col_idx - 1].dropna().astype(str)
            sample_max = int(sample_series.str.len().max()) if not sample_series.empty else 0
            max_len = max(len(col_name), sample_max, 12)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 42)

        styled_headers = []
        for h in headers:
            cell = WriteOnlyCell(ws, value=ILLEGAL_CHARACTERS_RE.sub("", h))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            styled_headers.append(cell)
        ws.append(styled_headers)

        # Stream rows via itertuples(index=False, name=None) to avoid allocating a massive duplicate list in RAM
        for row in sub_df.itertuples(index=False, name=None):
            ws.append([_sanitize_excel_value(v) for v in row])

    if sheets_written == 0:
        ws = wb.create_sheet(title="Summary")
        ws.append(["No data available"])

    wb.save(output)
    gc.collect()
    return output.getvalue()


# --- Streamlit Sidebar & UI ---
st.title("📊 DHIS2 TB Tracker Data Exporter")
st.caption(
    "Workflow: `data_dhis` (Raw API) → `df_dhis2` (Uncoded Cache) → Page Split "
    "(`df_registration`, `df_diagnosis`, `df_surveillance`, `df_contact`) → `df_combined` & `df_ygntbpro`"
)

with st.sidebar:
    st.header("⚙️ Server & Criteria Selection")
    with st.form("dhis2_form"):
        web_url = st.text_input("DHIS2 Base URL", value="https://mbdrtraining.moh.gov.mm/train")
        username = st.text_input("Username", value="NTP_YTP_Central_S2")
        password = st.text_input("Password", type="password", value="District@1")

        st.divider()
        st.subheader("🏥 Registration Criteria")
        selected_townships = st.multiselect(
            "Select Townships (Org Units)",
            options=list(TOWNSHIP_OU_MAP.keys()),
            default=list(TOWNSHIP_OU_MAP.keys()),
        )

        use_date_filter = st.checkbox("Filter Registration by Enrollment Date", value=False)
        d_col1, d_col2 = st.columns(2)
        with d_col1:
            start_date_input = st.date_input("Start Date", value=date(date.today().year, 1, 1))
        with d_col2:
            end_date_input = st.date_input("End Date", value=date.today())

        st.divider()
        st.subheader("📄 Export & Cache Options")
        selected_programs = st.multiselect(
            "Select Programs",
            options=list(PROGRAM_MAP.keys()),
            default=list(PROGRAM_MAP.keys()),
        )

        export_excel_default = st.checkbox(
            "Export Processed DataFrames to Excel (.xlsx)",
            value=True,
            help="Selected by default. Exports all processed DataFrames except df_dhis2.",
        )

        enable_raw_csv = st.checkbox(
            "Optional CSV Export for Uncoded Source (df_dhis2)",
            value=False,
            help="Generates a UTF-8 CSV download for df_dhis2.",
        )

        force_refresh = st.checkbox(
            "Force Re-Download from DHIS2 Server",
            value=False,
            help="Clears cached metadata and df_dhis2, pulling fresh data from the server.",
        )

        submitted = st.form_submit_button(
            "🚀 Extract, Process & Prepare Export", type="primary", use_container_width=True
        )

if submitted:
    prog_ids = tuple(dict.fromkeys([REGISTRATION_PROGRAM_ID] + [PROGRAM_MAP[p] for p in selected_programs]))
    selected_ou_ids: list[str] = []
    for township in selected_townships:
        selected_ou_ids.extend(TOWNSHIP_OU_MAP[township])

    if not selected_ou_ids:
        st.warning("Please select at least one Township for `df_registration`.")
    elif use_date_filter and start_date_input > end_date_input:
        st.error("Start Date cannot be later than End Date.")
    else:
        if force_refresh:
            fetch_cached_metadata.clear()
            get_cached_source_df_dhis2.clear()
            st.session_state.pop("dhis2_results", None)
            st.session_state.pop("_proc_cache_key", None)

        start_str = start_date_input.strftime("%Y-%m-%d") if use_date_filter else None
        end_str = end_date_input.strftime("%Y-%m-%d") if use_date_filter else None

        _proc_cache_key = hashlib.md5(
            f"{normalize_base_url(web_url)}|{username}|{prog_ids}|{sorted(selected_ou_ids)}|"
            f"{sorted(selected_townships)}|{start_str}|{end_str}".encode()
        ).hexdigest()

        with st.status("Executing DHIS2 Extraction, Uncoding & Page Processing...", expanded=True) as status:
            try:
                status.write("🔄 Step 1 & 2: Extracting `data_dhis` and building uncoded cache `df_dhis2`...")
                df_dhis2, meta_map, raw_tei_count = get_cached_source_df_dhis2(
                    web=web_url,
                    username=username,
                    password=password,
                    prog_ids_tuple=prog_ids,
                )
            except Exception as exc:
                status.update(label="Connection or authentication error.", state="error", expanded=True)
                st.error(f"Failed to extract data from DHIS2: {exc}")
                df_dhis2 = pd.DataFrame()
                meta_map = {}
                raw_tei_count = 0

            if df_dhis2.empty:
                # Evict empty DataFrame from cache so next submit retries cleanly
                get_cached_source_df_dhis2.clear()
                status.update(label="Extraction returned 0 records.", state="error", expanded=True)
                st.error("No data extracted from DHIS2 server. Check your URL, credentials, or user permissions.")
            else:
                status.write(
                    f"✅ Uncoded source `df_dhis2` ready: **{len(df_dhis2):,}** enrollment rows "
                    f"(from **{raw_tei_count:,}** raw `data_dhis` instances)."
                )

                if (
                    st.session_state.get("_proc_cache_key") == _proc_cache_key
                    and "dhis2_results" in st.session_state
                ):
                    # Ensure requested exports exist if user toggled checkboxes without changing filters
                    res_existing = st.session_state["dhis2_results"]
                    if export_excel_default and res_existing.get("excel_bytes") is None:
                        status.write("📦 Step 5: Streaming processed DataFrames to Excel (`.xlsx`)...")
                        res_existing["excel_bytes"] = convert_df_to_excel_bytes(
                            build_excel_sheets_dict(res_existing)
                        )
                    if enable_raw_csv and res_existing.get("csv_bytes") is None:
                        status.write("📄 Preparing optional UTF-8 CSV for `df_dhis2`...")
                        res_existing["csv_bytes"] = reorder_columns(df_dhis2).to_csv(index=False).encode("utf-8-sig")
                    status.update(
                        label="✅ Filters unchanged — using cached processed results!",
                        state="complete",
                        expanded=False,
                    )
                else:
                    status.write(
                        "⚙️ Step 3 & 4: Splitting into `df_registration`, `df_diagnosis`, "
                        "`df_surveillance`, `df_contact` and building `df_combined` & `df_ygntbpro`..."
                    )

                    processed = split_and_process_dataframes(
                        df_dhis2=df_dhis2,
                        selected_ou_ids=selected_ou_ids,
                        selected_township_names=selected_townships,
                        start_date_str=start_str,
                        end_date_str=end_str,
                        meta_map=meta_map,
                    )

                    excel_bytes = None
                    if export_excel_default:
                        status.write("📦 Step 5: Streaming processed DataFrames to Excel (`.xlsx`)...")
                        excel_bytes = convert_df_to_excel_bytes(build_excel_sheets_dict(processed))

                    csv_bytes = None
                    if enable_raw_csv:
                        status.write("📄 Preparing optional UTF-8 CSV for `df_dhis2`...")
                        csv_bytes = reorder_columns(df_dhis2).to_csv(index=False).encode("utf-8-sig")

                    st.session_state["dhis2_results"] = {
                        **processed,
                        "source_row_count": len(df_dhis2),
                        "excel_bytes": excel_bytes,
                        "csv_bytes": csv_bytes,
                        "query_params": (web_url, username, password, prog_ids),
                    }
                    st.session_state["_proc_cache_key"] = _proc_cache_key
                    status.update(label="✅ Processing Complete!", state="complete", expanded=False)

# --- Display Summary, Downloads & DataFrame Previews ---
if "dhis2_results" in st.session_state:
    res = st.session_state["dhis2_results"]
    df_registration = res["df_registration"]
    df_diagnosis = res["df_diagnosis"]
    df_surveillance = res["df_surveillance"]
    df_contact = res["df_contact"]
    df_combined = res["df_combined"]
    df_ygntbpro = res["df_ygntbpro"]
    df_summary = res["df_summary"]
    source_row_count = res.get("source_row_count", 0)

    # 1. Top-Level Metrics for Each DataFrame
    st.subheader("📈 Data Summary & Record Counts by DataFrame")
    c1, c2, c3, c4, c5, c6, c7 = st.columns(7)
    c1.metric("df_dhis2 (Source)", f"{source_row_count:,}")
    c2.metric("df_registration", f"{len(df_registration):,}")
    c3.metric("df_diagnosis", f"{len(df_diagnosis):,}")
    c4.metric("df_surveillance", f"{len(df_surveillance):,}")
    c5.metric("df_contact", f"{len(df_contact):,}")
    c6.metric("df_combined", f"{len(df_combined):,}")
    c7.metric("df_ygntbpro", f"{len(df_ygntbpro):,}")

    st.dataframe(df_summary, use_container_width=True, hide_index=True)

    # 2. Export & Download Buttons
    st.subheader("💾 Export & Download")
    dl_col1, dl_col2 = st.columns(2)
    with dl_col1:
        if res.get("excel_bytes") is not None:
            st.download_button(
                label="📥 Download Processed Excel Workbook (.xlsx)",
                data=res["excel_bytes"],
                file_name="DHIS2_Processed_TB_Tracker.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )
        else:
            if st.button("📦 Generate Excel Workbook (.xlsx) Now", type="primary", use_container_width=True):
                res["excel_bytes"] = convert_df_to_excel_bytes(build_excel_sheets_dict(res))
                st.rerun()

    with dl_col2:
        if res.get("csv_bytes") is not None:
            st.download_button(
                label=f"📄 Download Uncoded Source df_dhis2 (.csv) [{source_row_count:,} rows]",
                data=res["csv_bytes"],
                file_name="df_dhis2_uncoded_source.csv",
                mime="text/csv",
                use_container_width=True,
            )
        else:
            if st.button("⚙️ Generate Optional CSV for df_dhis2", use_container_width=True):
                q_web, q_user, q_pass, q_progs = res.get(
                    "query_params",
                    (web_url, username, password, tuple(PROGRAM_MAP.values())),
                )
                cached_df_dhis2, _, _ = get_cached_source_df_dhis2(q_web, q_user, q_pass, q_progs)
                res["csv_bytes"] = reorder_columns(cached_df_dhis2).to_csv(index=False).encode("utf-8-sig")
                st.rerun()

    st.divider()

    # 3. Interactive Preview Tabs for Each DataFrame
    st.subheader("🔍 Preview DataFrames")
    _PREVIEW_ROWS = 300

    tabs = st.tabs([
        f"📋 df_ygntbpro ({len(df_ygntbpro):,})",
        f"🔗 df_combined ({len(df_combined):,})",
        f"🏥 df_registration ({len(df_registration):,})",
        f"🔬 df_diagnosis ({len(df_diagnosis):,})",
        f"💊 df_surveillance ({len(df_surveillance):,})",
        f"🛡️ df_contact ({len(df_contact):,})",
        f"🗄️ df_dhis2 (Uncoded Source) ({source_row_count:,})",
    ])

    with tabs[0]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_ygntbpro):,}** rows — download Excel for full data.")
        st.dataframe(df_ygntbpro.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[1]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_combined):,}** rows — download Excel for full data.")
        st.dataframe(df_combined.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[2]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_registration):,}** rows — download Excel for full data.")
        st.dataframe(df_registration.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[3]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_diagnosis):,}** rows — download Excel for full data.")
        st.dataframe(df_diagnosis.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[4]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_surveillance):,}** rows — download Excel for full data.")
        st.dataframe(df_surveillance.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[5]:
        st.caption(f"Previewing first {_PREVIEW_ROWS:,} of **{len(df_contact):,}** rows — download Excel for full data.")
        st.dataframe(df_contact.head(_PREVIEW_ROWS), use_container_width=True, height=400)
    with tabs[6]:
        st.caption(
            f"Previewing first {_PREVIEW_ROWS:,} of **{source_row_count:,}** rows of cached `df_dhis2`. "
            "Use the optional CSV export button above to download all rows."
        )
        q_web, q_user, q_pass, q_progs = res.get(
            "query_params",
            (web_url, username, password, tuple(PROGRAM_MAP.values())),
        )
        cached_df_dhis2, _, _ = get_cached_source_df_dhis2(q_web, q_user, q_pass, q_progs)
        st.dataframe(
            sanitize_df_for_display(reorder_columns(cached_df_dhis2.head(_PREVIEW_ROWS))),
            use_container_width=True,
            height=400,
        )